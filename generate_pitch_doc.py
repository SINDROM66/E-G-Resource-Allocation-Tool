from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent
OUTPUT = ROOT / "E&G_Resource_Allocation_Tool_Senior_Management_Pitch.docx"
SCREENSHOTS = ROOT / "screenshots"

GREEN = RGBColor(18, 58, 44)
GOLD = RGBColor(199, 154, 62)
INK = RGBColor(18, 32, 26)
MUTED = RGBColor(92, 91, 80)


def set_cell_shading(cell, fill):
    cell._tc.get_or_add_tcPr().append(
        __import__("docx").oxml.parse_xml(
            f'<w:shd {__import__("docx").oxml.ns.nsdecls("w")} w:fill="{fill}"/>'
        )
    )


def heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.style.font.color.rgb = GREEN
    return paragraph


def body(document, text, bold_lead=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    if bold_lead and text.startswith(bold_lead):
        paragraph.add_run(bold_lead).bold = True
        paragraph.add_run(text[len(bold_lead):])
    else:
        paragraph.add_run(text)
    return paragraph


def bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def add_screenshot(document, filename, caption):
    path = SCREENSHOTS / filename
    if path.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(6.25))
        caption_paragraph = document.add_paragraph(caption)
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.style.font.italic = True
        caption_paragraph.style.font.color.rgb = MUTED
        caption_paragraph.paragraph_format.space_after = Pt(10)


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = document.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = INK
for style_name in ("Heading 1", "Heading 2"):
    styles[style_name].font.name = "Aptos Display"
    styles[style_name].font.color.rgb = GREEN

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(2)
run = title.add_run("E&G Resource Allocation Tool")
run.bold = True
run.font.size = Pt(25)
run.font.color.rgb = GREEN

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(16)
run = subtitle.add_run("A management view of partner outreach, deployment, approvals, and results")
run.font.size = Pt(12)
run.font.color.rgb = GOLD

callout = document.add_table(rows=1, cols=1)
callout.alignment = WD_TABLE_ALIGNMENT.CENTER
cell = callout.cell(0, 0)
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
set_cell_shading(cell, "F3EBD8")
cell.text = "Pitch message: The tool turns scattered outreach activity into one visible operating picture, so E&G leadership can approve, deploy, monitor, and learn from field work with less delay and less manual follow-up."
cell.paragraphs[0].paragraph_format.space_after = Pt(0)

heading(document, "What the tool does", 1)
body(document, "The E&G Resource Allocation Tool is a practical deployment ledger for Enterprise & Growth. It connects partner sourcing, outreach planning, staff assignment, approval workflows, field-time monitoring, and results in one place.")
body(document, "It is designed to help Senior Management answer four questions quickly:")
for item in [
    "What outreach activity is planned, pending, or awaiting approval?",
    "Which partners and regions are producing the strongest onboarding results?",
    "Are the right staff available and assigned to the work?",
    "Where are per diem or field-time exceptions requiring action?",
]:
    bullet(document, item)

heading(document, "Management value", 1)
for item in [
    "Faster decisions: approval queues and urgent items are visible at a glance.",
    "Better accountability: each outreach has a partner, region, date, team, status, and outcome trail.",
    "Smarter deployment: staff skills and availability support more deliberate team matching.",
    "Stronger cost control: per diem approval is linked to the outreach before departure.",
    "Evidence-led improvement: partner conversion and outreach results show where effort is working.",
]:
    bullet(document, item)

heading(document, "Suggested 3-minute pitch flow", 1)
steps = [
    ("1. Start with the dashboard", "Show the approval queue, urgent per diem item, field-time exceptions, and staff availability."),
    ("2. Move to the partner registry", "Show how active partners are ranked by onboarding conversion and how pipeline leads can be tracked."),
    ("3. Open outreach planning", "Show the register of planned, pending, approved, and completed field activity."),
    ("4. Finish with Per Diem", "Show pre-departure approval and the field-time tracker that flags trips running beyond plan."),
    ("5. Close on governance", "Emphasise role-based access: Senior Management can manage approvals and users, while Field Staff see the work relevant to them."),
]
for lead, detail in steps:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.add_run(lead + ": ").bold = True
    paragraph.add_run(detail)

document.add_page_break()
heading(document, "Screens in the tool", 1)
body(document, "The following screens are representative views from the working application and can be used as visual prompts during the pitch.")

add_screenshot(document, "login.png", "1. Role-based access: Senior Manager and Field Staff entry points.")
add_screenshot(document, "dashboard.png", "2. Executive dashboard: approvals, urgent exceptions, staff availability, and sign-out control.")

document.add_page_break()
add_screenshot(document, "partners.png", "3. Partner Registry & Pipeline: active partners, conversion rates, regions, sectors, and sourcing leads.")
add_screenshot(document, "outreach-plan.png", "4. Plan an Outreach: one register for dates, audience, team assignment, status, and per diem state.")

document.add_page_break()
add_screenshot(document, "per-diem.png", "5. Per Diem: approval queue before departure and field-time tracking for trips beyond plan.")

heading(document, "Suggested closing statement", 1)
body(document, "This tool gives E&G a shared operational picture without adding another complicated reporting layer. It makes the next action visible: approve the trip, assign the team, follow up with the partner, or investigate an exception. The result is a more controlled deployment cycle and a clearer link between field activity and onboarding outcomes.")

heading(document, "Current scope", 1)
body(document, "The hosted version is a working prototype for demonstration and internal validation. It currently uses seeded data and browser/local API storage; a production rollout would add shared database hosting, organisation authentication, and server-side permissions.")

document.save(OUTPUT)
print(OUTPUT)